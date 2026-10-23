# Builds an import-test workbook in the template layout: 3 formus_ tables with primary name, primary key,
# N:1 lookups (to each other and to Account), an N:N, and one field of each common type.
import copy, sys, openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
proto = wb['Need (formus_Need)']
P = 'formus'
O = lambda n: str(75047 * 10000 + n)   # choice values with the publisher's choice value prefix

def f(ln, disp, typ, add='', req='None', rt='', cd='', desc=None, kind='Simple'):
    SCH = dict(testprojectid='TestProjectId', name='Name', description='Description', budget='Budget', startdate='StartDate', status='Status',
               accountid='AccountId', testtaskid='TestTaskId', projectid='ProjectId', resourceid='ResourceId', estimatedhours='EstimatedHours',
               priority='Priority', duedate='DueDate', complete='Complete', tags='Tags', testresourceid='TestResourceId', email='Email',
               dayrate='DayRate', active='Active')
    sch = P + '_' + SCH[ln[len(P) + 1:]]
    return [ln, sch, disp, typ, desc or disp, 'True', kind, req, add, '', '', rt, '', cd]

tables = [
  dict(tab='Project (formus_TestProject)', disp='Test Project', logical=P + '_testproject', schema=P + '_TestProject', area='Import test',
       purpose='Import test: a project with tasks', rows=[
    f(P + '_testprojectid', 'Test Project', 'Uniqueidentifier', desc='Primary key'),
    f(P + '_name', 'Name', 'Text', 'Format: Text\nMax length: 200', req='ApplicationRequired', desc='Primary name'),
    f(P + '_description', 'Description', 'Multiline Text', 'Format: TextArea\nMax length: 4000'),
    f(P + '_budget', 'Budget', 'Currency', 'Minimum value: 0\nMaximum value: 1000000000\nPrecision: 2'),
    f(P + '_startdate', 'Start date', 'DateTime', 'Format: DateOnly'),
    f(P + '_status', 'Project status', 'Choice', 'Options:\n%s: Planned\n%s: Active\n%s: Closed\nDefault: %s' % (O(0), O(1), O(2), O(0))),
    f(P + '_accountid', 'Client', 'Lookup', 'Targets:\naccount', rt='N:1', cd='Restrict', desc='Client account (relationship to an existing table)'),
    [P + '_testproject_testresource', '', 'Project resources', '', 'Many-to-many: resources on a project', 'True', '', '', 'Targets:\n' + P + '_testresource', '', '', 'N:N', '', ''],
  ]),
  dict(tab='Task (formus_TestTask)', disp='Test Task', logical=P + '_testtask', schema=P + '_TestTask', area='Import test',
       purpose='Import test: work items on a project', rows=[
    f(P + '_testtaskid', 'Test Task', 'Uniqueidentifier', desc='Primary key'),
    f(P + '_name', 'Name', 'Text', 'Format: Text\nMax length: 200', req='ApplicationRequired', desc='Primary name'),
    f(P + '_projectid', 'Project', 'Lookup', 'Targets:\n' + P + '_testproject', req='ApplicationRequired', rt='N:1', cd='Cascade', desc='Deleting a project deletes its tasks'),
    f(P + '_resourceid', 'Assigned to', 'Lookup', 'Targets:\n' + P + '_testresource', rt='N:1', cd='Remove Link'),
    f(P + '_estimatedhours', 'Estimated hours', 'Decimal', 'Minimum value: 0\nMaximum value: 1000\nPrecision: 2'),
    f(P + '_priority', 'Priority', 'Whole number', 'Minimum value: 1\nMaximum value: 5'),
    f(P + '_duedate', 'Due', 'DateTime', 'Format: DateAndTime'),
    f(P + '_complete', 'Complete', 'Two options', 'True: Yes\nFalse: No\nDefault Value: False'),
    f(P + '_tags', 'Tags', 'Choices', 'Options:\n%s: Urgent\n%s: Blocked\n%s: Review' % (O(10), O(11), O(12))),
  ]),
  dict(tab='Resource (formus_TestResource)', disp='Test Resource', logical=P + '_testresource', schema=P + '_TestResource', area='Import test',
       purpose='Import test: people or equipment assigned to tasks', rows=[
    f(P + '_testresourceid', 'Test Resource', 'Uniqueidentifier', desc='Primary key'),
    f(P + '_name', 'Name', 'Text', 'Format: Text\nMax length: 100', req='ApplicationRequired', desc='Primary name'),
    f(P + '_email', 'Email', 'Text', 'Format: Email\nMax length: 100'),
    f(P + '_dayrate', 'Day rate', 'Currency', 'Minimum value: 0\nMaximum value: 10000\nPrecision: 2'),
    f(P + '_active', 'Active', 'Two options', 'True: Active\nFalse: Inactive\nDefault Value: True'),
  ]),
]

def cp_style(dst, s):
    dst.font = copy.copy(s.font); dst.fill = copy.copy(s.fill); dst.border = copy.copy(s.border)
    dst.alignment = copy.copy(s.alignment); dst.number_format = s.number_format; dst.protection = copy.copy(s.protection)

for t in tables:
    ws = wb.copy_worksheet(proto); ws.title = t['tab']
    meta = [t['disp'], t['disp'] + 's', t['purpose'], t['schema'], t['logical'], '', 'True', 'UserOwned', 'True',
            P + '_name', t['logical'] + 'id', 'False', 'True', 'False']
    for i, v in enumerate(meta): ws.cell(1 + i, 2).value = v if v != '' else None
    # clear the example's field list and its Forms/Views block, then write the fields and move Forms/Views below them
    sec = {r: [(copy.copy(ws.cell(r, c)._style), ws.cell(r, c).value) for c in range(1, 15)] for r in (20, 21, 22, 24, 25, 26)}
    data_style = [copy.copy(proto.cell(17, c)._style) for c in range(1, 15)]
    for r in range(17, ws.max_row + 1):
        for c in range(1, 15): ws.cell(r, c).value = None
    for k, row in enumerate(t['rows']):
        for c, v in enumerate(row, 1):
            cell = ws.cell(17 + k, c); cell.value = v if v != '' else None; cell._style = copy.copy(data_style[c - 1])
    base = 17 + len(t['rows']) + 2
    for off, r in enumerate((20, 21, 22, None, 24, 25, 26)):
        if r is None: continue
        for c, (st, v) in enumerate(sec[r], 1):
            cell = ws.cell(base + off, c); cell._style = st; cell.value = v
    ws.cell(base + 2, 2).value = 'Main form'; ws.cell(base + 6, 2).value = 'Active ' + t['disp'] + 's'; ws.cell(base + 6, 3).value = 'Name, Created On'
    for r in range(base, 18, -1): pass
    ws.data_validations.dataValidation = []
    last = 16 + len(t['rows'])
    for col, f1 in (('L', '"N:1,N:N"'), ('N', '"Remove Link,Restrict,Cascade"')):
        dv = DataValidation(type='list', formula1=f1, allow_blank=True); dv.add('%s17:%s%d' % (col, col, last)); ws.add_data_validation(dv)
    ws.auto_filter.ref = 'A16:N%d' % last

wb.remove(proto)
td = wb['Table Dictionary']
td.cell(2, 3).value = 'Import test model: three custom tables to check the generated solution imports'
for k, t in enumerate(tables):
    for c, v in enumerate(['', t['area'], t['disp'], 'Custom', 'No', t['purpose'], '', 'Yes', 'Yes'], 1):
        cell = td.cell(4 + k, c); cell.value = v or None
        if k: cell._style = copy.copy(td.cell(4, c)._style)
wb.save(out)
print('saved', out)
