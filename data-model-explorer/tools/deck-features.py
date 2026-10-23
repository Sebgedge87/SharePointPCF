# Builds the detailed Data Model Explorer feature deck on the FormusPro template.
# python3 build_detail.py <template.pptx> <out.pptx> <images dir (the user guide's marked screenshots)>
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

TPL, OUT, IMG = sys.argv[1], sys.argv[2], sys.argv[3]
prs = Presentation(TPL)
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)

L = {i + 1: l for i, l in enumerate(prs.slide_layouts)}
NAVY, PINK, PALE = MSO_THEME_COLOR.ACCENT_1, MSO_THEME_COLOR.ACCENT_2, MSO_THEME_COLOR.ACCENT_5
WHITE, DARK = MSO_THEME_COLOR.BACKGROUND_1, MSO_THEME_COLOR.TEXT_1
MUTED = RGBColor(0x3E, 0x55, 0x6B)
NEWC = RGBColor(0xE8, 0xA3, 0x17)


def new(layout, keep=()):
    s = prs.slides.add_slide(L[layout])
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx not in keep: ph._element.getparent().remove(ph._element)
    return s


def put(ph, lines, size=None, bold=None, color=None):
    tf = ph.text_frame; lines = [lines] if isinstance(lines, str) else lines
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line
        if size: r.font.size = Pt(size)
        if bold is not None: r.font.bold = bold
        if color is not None: r.font.color.theme_color = color
    return tf


def text(s, x, y, w, h, runs, size=14, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=6, name=None):
    """runs: list of paragraphs; a paragraph is a string or a list of (text, bold) pieces"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name: tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'): setattr(tf, m, 0)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        for piece, b in ([(para, bold)] if isinstance(para, str) else para):
            r = p.add_run(); r.text = piece; r.font.size = Pt(size); r.font.bold = b
            if isinstance(color, RGBColor): r.font.color.rgb = color
            else: r.font.color.theme_color = color
    return tb


def card(s, x, y, w, h, fill=PALE, name=None):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.adjustments[0] = 0.08; c.fill.solid(); c.fill.fore_color.theme_color = fill; c.line.fill.background(); c.shadow.inherit = False
    if name: c.name = name
    c.text_frame.text = ''
    return c


def shot(s, file, x, y, bw, bh, alt):
    """the screenshot, as large as fits the box, centred across it and from its top"""
    iw, ih = Image.open(f'{IMG}/{file}.png').size
    sc = min(bw / iw, bh / ih); w, h = iw * sc, ih * sc
    p = s.shapes.add_picture(f'{IMG}/{file}.png', Inches(x + (bw - w) / 2), Inches(y), Inches(w), Inches(h))
    p.line.color.theme_color = PALE; p.line.width = Pt(1)
    p._element.nvPicPr.cNvPr.set('descr', alt); p.name = 'Screenshot: ' + file
    return h


def points(s, x, y, w, items, gap):
    for i, (head, body) in enumerate(items):
        text(s, x, y + i * gap, w, 0.35, [head], size=17, color=PINK, bold=True, name='Point heading')
        text(s, x, y + i * gap + 0.34, w, gap - 0.4, [body], size=13, color=MUTED, name='Point text')


def bullets(ph, size=None):
    for p in ph.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr(); pPr.set('marL', '228600'); pPr.set('indent', '-228600')
        for old in pPr.findall(qn('a:buNone')): pPr.remove(old)
        bu = etree.SubElement(pPr, qn('a:buChar')); bu.set('char', '•')
        ln = etree.SubElement(pPr, qn('a:spcBef')); etree.SubElement(ln, qn('a:spcPts')).set('val', '500')
        pPr.remove(ln); pPr.insert(0, ln)


def notes(s, t): s.notes_slide.notes_text_frame.text = t


def table(s, rows, x, y, w, widths, size=13, h=None):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h or 0.4 * len(rows))).table
    for i, cw in enumerate(widths): t.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = t.cell(r, c); cell.fill.solid(); cell.fill.fore_color.theme_color = NAVY if r == 0 else (PALE if r % 2 == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.12); cell.margin_top = cell.margin_bottom = Inches(0.06); cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]; run = p.add_run(); run.text = v; run.font.size = Pt(size); run.font.bold = (r == 0 or c == 0)
            run.font.color.theme_color = WHITE if r == 0 else (NAVY if c == 0 else DARK)
    return t


def badge(s, x, y, label='NEW IN 1.9.3'):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.6), Inches(0.34))
    b.adjustments[0] = 0.5; b.fill.solid(); b.fill.fore_color.rgb = NEWC; b.line.fill.background(); b.shadow.inherit = False; b.name = 'New badge'
    tf = b.text_frame; tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = label; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = RGBColor(0x1D, 0x2B, 0x36)


def screen(layout, title, file, alt, items, note, gap=1.25, new_badge=False):
    """a screenshot with points beside it, on one of the template's two framed layouts"""
    s = new(layout, keep=(13,))
    put(s.placeholders[13], title)
    if layout == 12: x0, tx, tw = 1.63, 9.15, 3.6      # navy bar on the left
    else: x0, tx, tw = 0.52, 8.0, 3.7                  # navy strip on the right
    shot(s, file, x0, 1.65, 7.2, 5.15, alt)
    points(s, tx, 1.65, tw, items, gap)
    if new_badge: badge(s, 6.35, 0.72)
    notes(s, note)
    return s


SECTIONS = [('Getting started', 'Opening workbooks, where they are kept, and the template'),
            ('Reviewing the model', 'The Overview, browsing and searching, Checks, and comparing with Dataverse'),
            ('The ERD', 'Choosing tables, reading the lines, hiding relationships, and exports'),
            ('Editing', 'Values, keys and table details, adding fields and relationships, and the change list'),
            ('Versions and saving', 'Loading a newer workbook, exporting, and saving back safely'),
            ('Building in Dataverse', 'The solution builder, what it creates, importing, and what to do after')]


def divider(n):
    h, sub = SECTIONS[n - 1]
    s = new(2, keep=(14,))
    ph = s.placeholders[14]; ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(3.05), Inches(11.5), Inches(1.15)
    put(ph, h); ph.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT; ph.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    text(s, 0.6, 1.75, 3, 1.2, [f'{n:02d}'], size=72, color=PINK, bold=True, name='Section number')
    text(s, 0.6, 4.6, 11.5, 0.8, [sub], size=20, color=MUTED, name='Section summary')
    notes(s, f'Section {n}: {h}. {sub}.')


# ---------------------------------------------------------------- cover, contents, principles
s = new(3, keep=(14,))
put(s.placeholders[14], 'Data Model Explorer')
text(s, 0.6, 3.6, 12.13, 0.6, ['Feature guide · version 1.9.3'], size=22, color=MUTED, align=PP_ALIGN.CENTER, name='Subtitle')
notes(s, 'A detailed tour of every feature of the Data Model Explorer, section by section, with screenshots from the three-table import test model. The shorter overview deck covers the why; this one covers the how.')

s = new(6, keep=(13,))
put(s.placeholders[13], 'Contents')
for i, (h, sub) in enumerate(SECTIONS):
    x, y = 0.54 + (i % 3) * 4.15, 1.75 + (i // 3) * 2.45
    card(s, x, y, 3.85, 2.15, name='Section card')
    text(s, x + 0.3, y + 0.2, 1.0, 0.6, [f'{i + 1:02d}'], size=28, color=PINK, bold=True)
    text(s, x + 0.3, y + 0.8, 3.25, 0.4, [h], size=19, color=NAVY, bold=True)
    text(s, x + 0.3, y + 1.22, 3.25, 0.85, [sub], size=13, color=MUTED)
notes(s, 'Six sections, following the order you use the page in: get the workbook in, review it, draw it, edit it, keep versions and save, then build it in Dataverse.')

s = new(17, keep=(13, 17, 18, 19, 20, 16, 14, 15, 12))
put(s.placeholders[13], 'How it is built')
for hdr, body, h, b in [(17, 16, 'One file', 'Data_Model_Explorer.html holds everything: the app, the libraries and the empty template. Open it in Chrome or Edge.'),
                        (18, 14, 'Offline and private', 'Nothing is installed or uploaded. Workbooks and edits stay in this browser profile only.'),
                        (19, 15, 'Template-driven', 'Columns, table details and dropdowns come from the FormusPro template, which you can replace.'),
                        (20, 12, 'Safe with your files', 'Only the cells you change are written, and the file is read back and checked before it is saved.')]:
    put(s.placeholders[hdr], h); put(s.placeholders[body], b, size=16)
notes(s, 'Four design choices run through every feature. Because the HTML file holds only the empty template, it can be shared freely: anyone who opens it starts with nothing of yours.')

# ---------------------------------------------------------------- 1 getting started
divider(1)
screen(15, 'Opening a workbook', '01-open', 'The page header with Open workbook… marked',
       [('Open or drop', 'Click Open workbook… in the header, or drag the .xlsx onto the page'),
        ('Metadata exports converted', 'A Metadata Document Generator export is converted to the template layout, with progress shown'),
        ('Lands on the Overview', 'Tables are listed in the sidebar, grouped by area from the Table Dictionary'),
        ('Chrome or Edge', 'Needed to save back to the original file; everything else works in any modern browser')],
       'The first time, the page shows the empty template and says so. A workbook already in the template layout opens straight away; a Metadata Document Generator export is converted first and keeps the export’s name.', gap=1.25)
screen(12, 'Several workbooks, kept apart', '03-workbooks', 'The workbook selector menu listing workbooks, with Manage files… marked',
       [('Newer or different?', 'Opening a second file asks: a newer version carries your edits over; a different workbook gets its own'),
        ('Switch any time', 'The workbook selector in the header lists every workbook kept in this browser'),
        ('Manage files…', 'Shows them all, and Remove all… clears them'),
        ('Where it is kept', 'Workbooks in IndexedDB, edits and layouts in local storage. Saved in the header confirms it')],
       'Edits are kept per workbook, so two clients’ models never mix. Nothing is kept in the HTML file itself: send it to someone and they start with the empty template.', gap=1.25)
screen(15, 'The template behind the page', '15-template', 'The Template panel with Download blank template, Use a different template… and Reset to bundled marked',
       [('Built in', 'The current FormusPro template sets the columns, table details and dropdown options'),
        ('Check a workbook', 'More ▾ → Template lists tables, columns and details that differ from it'),
        ('Older workbooks upgraded', 'Missing columns and details are added in the page, marked added, and written when filled in'),
        ('Use your own', 'Download blank template, change it in Excel, then Use a different template…')],
       'Opening a workbook made from an older template adds the columns and table details it lacks, such as Relationship Type, Cascade Delete and Primary Name Column. A pink count on the Template button shows how many tables differ. Reset to bundled goes back to the built-in template.', gap=1.25)

s = new(6, keep=(13,))
put(s.placeholders[13], 'What a template change does')
table(s, [('Change in the template', 'Effect in the page'),
          ('A new column in the header row', 'Read, shown, filterable and editable; added to workbooks that lack it'),
          ('Text in the hint row above a column', 'Shown as the column’s tooltip and written with the column'),
          ('An Excel data-validation list', 'Its options become the column’s dropdown'),
          ('A new label row in the table details', 'Shown in ✎ Table details and added to tables that lack it'),
          ('Fonts, colours, widths', 'Used by New version in template format')],
      0.54, 1.75, 12.2, (4.2, 8.0), size=15, h=3.6)
text(s, 0.54, 5.75, 12.2, 0.8, ['Keep the layout when editing the template: the Table Dictionary, and one example table tab with its detail rows, hint row and header row starting Logical Name, Schema Name. A new column needs a code change only if the solution builder or a check should use it.'], size=14, color=MUTED)
notes(s, 'The template is the contract between the workbook and the page. Most changes take effect without any change to the page: reopen the workbook or reload the page after switching templates.')

# ---------------------------------------------------------------- 2 reviewing
divider(2)
screen(12, 'The Overview', '02-overview', 'The Overview with a count card, Open all checks, Build solution… and Check against Dataverse… marked',
       [('Counts that open lists', 'Tables, fields, lookups, related pairs, unresolved lookups, errors, warnings, your edits'),
        ('Completeness', 'Descriptions, areas and purposes, and the least described tables'),
        ('By area', 'Tables, fields, how much is described and issues, per area'),
        ('Shortcuts', 'Open all checks, Build solution… and Check against Dataverse…')],
       'The landing page once a workbook is open. When the page reopens a workbook from last time, the Overview says which one, so you can switch if it isn’t the one you want. Latest changes shows your recent edits, or the differences from the previous version.', gap=1.25)
screen(15, 'Browsing and searching', '06-table', 'A table page with the search box, Add field, Columns, a column filter and Table details marked',
       [('Every tab is a page', 'Details along the top, counts by field type, then the field list'),
        ('Search and filter', 'The header search (shortcut /), a funnel on every column, Edited only, Area'),
        ('All fields', 'Every field of every table in one searchable list'),
        ('Your view', 'Columns ▾ shows or hides columns; expand or collapse long rows')],
       'Schema Name, Type and Relationship Name are hidden by default and can be shown from Columns. Active filters show above the grid with a × to clear each one. Clicking a field-type count filters to that type.', gap=1.25)
screen(12, 'Checks', '04-checks', 'The Checks page with the Errors, Warnings, Info and Export (CSV) buttons marked',
       [('Always current', 'Runs over the whole model and updates as you edit'),
        ('Straight to the field', 'Each finding names the table and field and opens it when clicked'),
        ('Filter', 'Errors, Warnings and Info with counts, plus Area and search'),
        ('Share', 'Export (CSV) downloads the list')],
       'Clear the errors before building a solution: a table with an error is not built. Checks starting “Dataverse” appear only after loading Dataverse metadata.', gap=1.25)

s = new(6, keep=(13,))
put(s.placeholders[13], 'What Checks look for')
table(s, [('Severity', 'Check', 'How to fix it'),
          ('Error', 'No primary name column, or not in the field list, or not single-line text', 'Set as primary name on a Text field'),
          ('Error', 'Primary key is not a unique identifier', 'Keep the table name + id; use an alternate key'),
          ('Error', 'Missing logical name; duplicate field or tab; relationship name used twice', 'Fill in, remove or rename'),
          ('Warning', 'No attribute type; lookup without targets or to a table not in the workbook', 'Fill in Attribute Type or Targets:'),
          ('Warning', 'Relationship type or cascade not valid; N:N with an unknown table', 'N:1 on lookups, N:N on relationship rows'),
          ('Warning', 'Table Dictionary and tabs don’t match', 'Match the dictionary to the tabs'),
          ('Info', 'Relationship type not recorded; calculated or rollup field to build by hand', 'Set N:1 once confirmed; nothing to fix'),
          ('Info', 'Logical name format; custom field without prefix; no primary key listed', 'Usually Microsoft’s own fields')],
      0.54, 1.65, 12.2, (1.35, 6.75, 4.1), size=12.5, h=4.9)
notes(s, 'The full list, with every check, is in the user guide. Errors block the build for that table, warnings are worth fixing before you build, and info findings are for review.')

screen(15, 'Comparing with Dataverse', '16-dataverse', 'The Dataverse check panel with the environment URL, the generated links and Choose JSON files… marked',
       [('Enter the environment', 'The page builds one metadata link per 20 tables'),
        ('Open and save', 'Open each link while signed in and save it (Ctrl+S) as .json'),
        ('Load the files', 'Choose JSON files… or drop them; add more any time'),
        ('See the differences', 'Fields, types, required levels, relationships and keys, in Checks and on each table')],
       'The page stays offline: you open the links yourself in a browser signed in to Dataverse. Findings include fields and relationships missing from Dataverse, types or required levels that differ, and custom Dataverse fields or relationships the workbook doesn’t document. Remove imported metadata forgets it.', gap=1.25)

# ---------------------------------------------------------------- 3 ERD
divider(3)
screen(12, 'Drawing the ERD', '05-erd', 'The ERD with All, Keys + lookups and draw.io marked',
       [('Pick tables', 'Tick them, filter and Select shown, or Custom only, All, None, or by area'),
        ('Core tables', '★ lays the diagram out around a table and brings its neighbours'),
        ('What boxes show', 'Names only, Keys + lookups or All fields; left to right or top to bottom'),
        ('Your layout', 'Drag tables; layouts are remembered. Opens at screen size; Fit fits the panel')],
       'The diagram is generated from the lookup fields, so it always matches the workbook. With many tables, draw a few dozen at a time. The page asks before drawing more than 80 tables.', gap=1.25)

s = new(18, keep=(13, 18, 14, 20, 19))
put(s.placeholders[13], 'Reading the ERD')
put(s.placeholders[18], 'The lines and boxes')
put(s.placeholders[14], ['Ends (one or many) show only where the relationship type is recorded in the workbook.',
                         'A relationship with no recorded type is a plain grey line, listed in Checks for review.',
                         'N:N relationships are dashed.',
                         'Custom tables have a pink border; standard and referenced tables a plain one. Or colour by area.'], size=15)
bullets(s.placeholders[14])
put(s.placeholders[20], 'Options')
put(s.placeholders[19], ['Selecting a table also selects the tables it relates to', 'Show tables that selected tables look up to',
                         'Hide tables with no relationships', 'Also show tables not in the workbook', 'Merge lines between the same two tables',
                         'Hide user, team and business unit', 'Show N:N relationships from Dataverse'], size=15)
bullets(s.placeholders[19])
notes(s, 'The diagram never guesses a relationship type: an unrecorded type is drawn neutrally and flagged, so the ERD can be trusted as a picture of what is written down.')

screen(15, 'Hiding relationships', '18-erd-hide', 'The ERD with Test Project’s relationship list open: the count button, an unticked N:N relationship, and show all marked',
       [('Every relationship listed', 'Click the count beside a table: its lookups →, lookups to it ←, and N:N ↔'),
        ('Untick to hide', 'Gone from the diagram and every export; a table it alone brought in drops out'),
        ('Per table, at a glance', 'Hide all or Show all per table; the button reads, say, 1 of 3 hidden'),
        ('Nothing moves', 'Tables keep their places; show all brings everything back; the choice is remembered')],
       'Useful for presentation diagrams: hide the audit lookups or a busy N:N without removing the tables. The ERD info line counts how many relationships you have hidden in the current diagram.', gap=1.25, new_badge=True)

s = new(6, keep=(13,))
put(s.placeholders[13], 'Taking the diagram with you')
for i, (h, b) in enumerate([('draw.io', 'An editable diagram with your layout, ready to keep working on in draw.io'),
                            ('SVG', 'Sharp at any size, for documents and wikis'),
                            ('PNG', 'A picture, up to 8,000 pixels across'),
                            ('.mmd', 'The Mermaid source as a file'),
                            ('Copy Mermaid', 'The source on the clipboard, for Markdown, Azure DevOps or GitHub')]):
    x = 0.54 + i * 2.48
    card(s, x, 1.85, 2.28, 2.6, name='Export card')
    text(s, x + 0.22, 2.05, 1.9, 0.45, [h], size=18, color=NAVY, bold=True)
    text(s, x + 0.22, 2.6, 1.9, 1.7, [b], size=14, color=MUTED)
text(s, 0.54, 4.85, 12.2, 0.9, ['Exports follow what you see: the tables you picked, the attributes shown, the relationships you hid, and your layout. Source opens the read-only Mermaid text beside the diagram.'], size=15, color=MUTED)
notes(s, 'All five exports are generated in the browser. The draw.io file keeps positions and colours, so it is the best starting point for a diagram you will keep editing by hand.')

# ---------------------------------------------------------------- 4 editing
divider(4)
screen(12, 'Editing values', '07-edit', 'A Required Level cell being edited with a dropdown of values',
       [('Double-click a cell', 'Or click a row and edit in its panel. Enter saves, Esc cancels'),
        ('Dropdowns', 'Fixed-value columns edit as a dropdown; + New value… adds one'),
        ('Visible edits', 'Edited cells are highlighted; the panel shows the original with revert'),
        ('Keyboard', 'Shift+Enter for a new line; j and k move between rows in the panel')],
       'Dropdown columns include Attribute Type, Required Level, Relationship Type and Cascade Delete. Their options come from the template’s data-validation lists. Every edit is stored in the browser as you make it.', gap=1.25)
screen(15, 'Keys and table details', '08-keys', 'A field panel with Set as primary key, Set as primary name and an editable Description marked',
       [('Shown in the header', 'Primary name and primary key; an unrecorded key shows as assumed'),
        ('Tagged in the list', 'The two fields carry Primary key and Primary name tags'),
        ('Set on any field', 'Set as primary key or Set as primary name in the field’s panel'),
        ('✎ Table details', 'Ownership, notes, activities and every other detail row, including missing ones')],
       'Dataverse primary keys are always unique identifiers created with the table, so marking a text or number field as the key is flagged in Checks; use an alternate key to keep other values unique. The primary name must be single-line text.', gap=1.25)
screen(12, 'Adding fields and relationships', '09-addfield', 'The Add field dialog with the kind, display name, field type and Add button marked',
       [('A field or an N:N', 'Choose the kind, then type the display name'),
        ('Names follow', 'Actual hours becomes formus_actualhours and formus_ActualHours'),
        ('Lookups and N:N', 'Pick the related table: Targets and the relationship type are filled in'),
        ('Then the details', 'The new row is tagged ADDED and opens for options, length, cascade')],
       'A name already in the table is refused. Added rows go at the end of the field list and are written into the empty rows after it when saving. On other sheets, such as the Table Dictionary, + Add row adds a blank row. Remove row takes an added row out again.', gap=1.25)
screen(15, 'Reviewing and undoing changes', '10-changes', 'The Changes panel with the Changes button, remove, revert and Revert all marked',
       [('Everything in one list', 'Rows added, table details changed, and each cell as original against new'),
        ('Go or revert', 'Jump to any edit, or undo it on its own'),
        ('Revert all', 'Undo everything at once'),
        ('Keep a record', 'Export log (CSV) downloads the list')],
       'Changes is also where a newer version’s conflicts are settled: values the new version changed that you had also edited are listed under Needs your attention.', gap=1.25)

# ---------------------------------------------------------------- 5 versions and saving
divider(5)
screen(12, 'Loading a newer version', '13-loadversion', 'The dialog asking whether the file is a newer version or a different workbook',
       [('Every difference', 'Tables, fields, values, columns, details, forms, views and lookup targets'),
        ('Edits carry over', 'Conflicts go to Needs your attention: Keep mine or Use new version'),
        ('Added rows too', 'If the new version already has the field, you choose which to keep'),
        ('Share the differences', 'Version differences (CSV) in the Export menu')],
       'Click Load version… or drop the new file and choose newer version. Tables are matched by tab name and fields by logical name, so row order doesn’t matter. An edit to a field the new version no longer has can be discarded in Changes.', gap=1.25)

s = new(12, keep=(13,))
put(s.placeholders[13], 'The Export menu')
shot(s, '11-export', 1.63, 1.6, 3.3, 5.3, 'The Export menu with Download an updated copy, New version in template format and Dataverse solution marked')
table(s, [('Export', 'What you get'),
          ('Save back to the original file', 'Your edits written into the file on disk (Chrome or Edge, opened with the button)'),
          ('Download an updated copy', 'The same file and formatting with your edits written in'),
          ('New version in template format', 'The whole workbook rebuilt and styled by the template, with a Change log sheet'),
          ('Full workbook, plain values', 'Values only, no styling'),
          ('Current view as CSV', 'The rows you can see, with visible or all columns'),
          ('Change log · Version differences', 'Your edits, or the differences between versions, as CSV')],
      5.3, 1.6, 7.45, (2.75, 4.7), size=12.5, h=4.6)
notes(s, 'Choose Save back or Download an updated copy to keep the client’s file exactly as it is apart from your edits; choose New version in template format to bring an old workbook fully up to the current template.')

s = new(20, keep=(13, 20, 19))
put(s.placeholders[13], 'Saving back safely')
put(s.placeholders[20], 'What is guaranteed')
put(s.placeholders[19], ['Only the cells you changed are rewritten; formatting and validations stay byte for byte',
                         'Added rows go into the empty rows after the field list, new details above it',
                         'If anything can’t be written, nothing is saved and your edits are kept'], size=15)
bullets(s.placeholders[19])
for i, (h, b) in enumerate([('Name the file', 'Model v1.2 is suggested as Model v1.3'), ('Write your edits', 'Into the original or a copy'),
                            ('Read it back', 'Every edit checked in the saved file'), ('Saved', 'Or the reason it was stopped')]):
    y = 1.85 + i * 1.35
    card(s, 7.3, y, 5.4, 0.95, name='Step')
    text(s, 7.6, y + 0.12, 4.9, 0.4, [h], size=18, color=NAVY, bold=True)
    text(s, 7.6, y + 0.5, 4.9, 0.35, [b], size=14, color=MUTED)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.85), Inches(y + 0.98), Inches(0.3), Inches(0.34))
        a.fill.solid(); a.fill.fore_color.theme_color = PINK; a.line.fill.background(); a.name = 'Arrow'
notes(s, 'The usual reason a save is stopped is that there is no empty row for an added field before the Forms section. Then use New version in template format, or add empty rows in Excel and reload. Every workbook download asks for a file name first, suggesting the next version number.')

# ---------------------------------------------------------------- 6 building
divider(6)
screen(12, 'Generating the solution', '17-solution', 'The Build a Dataverse solution dialog with publisher name, prefix, choice value prefix, Will be built and Generate marked',
       [('Solution and version', 'Unique and display name; raise the four-part version for each import'),
        ('Publisher', 'The unique name must exactly match the environment’s, or a second is created'),
        ('Prefixes', 'Publisher prefix and choice value prefix, suggested from the workbook'),
        ('Will be built', 'What is new, added to standard tables, to build by hand or left out')],
       'Export → Dataverse solution (.zip)… or Build solution… on the Overview. Tables and fields starting with the prefix are built. Generate downloads one zip named after the solution and version. Clear the errors in Checks first.', gap=1.25)

s = new(18, keep=(13, 18, 14, 20, 19))
put(s.placeholders[13], 'What the solution contains')
put(s.placeholders[18], 'Built for you')
put(s.placeholders[14], ['New tables with primary name, key, ownership, notes and audit',
                         'Your fields on standard tables, leaving the rest of them alone',
                         'A relationship per lookup and customer lookup, with cascade; every N:N',
                         'Lengths, formats, ranges, precision, choices, defaults, required level, field security'], size=15)
bullets(s.placeholders[14])
put(s.placeholders[20], 'Listed in the report instead')
put(s.placeholders[19], ['Calculated and rollup fields (with a workbook to build them)', 'Activity tables and lookups to several tables',
                         'Image and file fields; choices with no options', 'Tables with no primary name, or a key that isn’t a GUID', 'Forms and views'], size=15)
bullets(s.placeholders[19])
notes(s, 'The solution XML matches what Dataverse itself exports, checked element by element against exported solutions. Anything not built is listed in build-report.csv with the reason.')

s = new(6, keep=(13,))
put(s.placeholders[13], 'One download')
table(s, [('File', 'What it is'),
          ('<SolutionName>_<version>.zip', 'The unmanaged solution'),
          ('import.ps1', 'Imports it with the Power Platform CLI, installing it if needed'),
          ('build-report.csv', 'Every table, field and relationship: included, not included (and why), or to build by hand'),
          ('calculated-and-rollup-fields.xlsx', 'The fields to create by hand, per table, with a column for the formula'),
          ('README.txt', 'The import steps')],
      0.54, 1.75, 12.2, (4.3, 7.9), size=15, h=3.4)
text(s, 0.54, 5.55, 12.2, 0.8, ['Export → Import script (import.ps1) downloads the script on its own, for a solution zip you already have.'], size=15, color=MUTED)
notes(s, 'Unzip the download into a folder and run the script from there. The script imports the one solution zip in its folder, or the one given with -SolutionPath.')

s = new(20, keep=(13, 20, 19))
put(s.placeholders[13], 'Importing with import.ps1')
put(s.placeholders[20], 'Options')
put(s.placeholders[19], ['-EnvironmentUrl: the environment to import into', '-Check: run Microsoft’s solution checker first',
                         '-SolutionPath: import a different zip', 'Blocked? Run Unblock-File .\\import.ps1 once'], size=15)
bullets(s.placeholders[19])
for i, (h, b) in enumerate([('Unzip', 'Into a folder, then open PowerShell there'), ('Run the script', '.\\import.ps1 -EnvironmentUrl https://…'),
                            ('Sign in', 'pac is installed if missing; a browser opens'), ('Imported and published', 'Into development first')]):
    y = 1.85 + i * 1.35
    card(s, 7.3, y, 5.4, 0.95, name='Step')
    text(s, 7.6, y + 0.12, 4.9, 0.4, [h], size=18, color=NAVY, bold=True)
    text(s, 7.6, y + 0.5, 4.9, 0.35, [b], size=14, color=MUTED)
    if i < 3:
        a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.85), Inches(y + 0.98), Inches(0.3), Inches(0.34))
        a.fill.solid(); a.fill.fore_color.theme_color = PINK; a.line.fill.background(); a.name = 'Arrow'
notes(s, 'You need PowerShell 7 or Windows PowerShell 5.1, and the System Customizer or System Administrator role. A large model can take several minutes to import. If the import fails, the message says why and nothing is changed: fix the workbook, raise the version and generate again.')

s = new(6, keep=(13,))
put(s.placeholders[13], 'After the import')
for i, (h, b) in enumerate([('Check the report', 'Deal with anything marked Not included in build-report.csv'),
                            ('Build by hand', 'Create the calculated and rollup fields with their formulas'),
                            ('Turn on activities', 'For tables with Has Activities set'),
                            ('Forms and views', 'Lay them out in Power Apps'),
                            ('Compare', 'Optionally, run the Dataverse check to confirm everything landed'),
                            ('Promote', 'Export from development as managed for test and production')]):
    x, y = 0.54 + (i % 3) * 4.15, 1.75 + (i // 3) * 2.45
    card(s, x, y, 3.85, 2.15, name='Step card')
    text(s, x + 0.3, y + 0.2, 0.8, 0.6, [str(i + 1)], size=28, color=PINK, bold=True)
    text(s, x + 0.3, y + 0.8, 3.25, 0.4, [h], size=19, color=NAVY, bold=True)
    text(s, x + 0.3, y + 1.22, 3.25, 0.85, [b], size=14, color=MUTED)
notes(s, 'Has Activities is not a setting in the solution file, so tables that need it are listed as a step to do by hand after the import.')

s = new(6, keep=(13,))
put(s.placeholders[13], 'Common questions')
table(s, [('Problem or question', 'What to do'),
          ('A table shows that isn’t in my workbook', 'Pick yours in the workbook selector, or Manage files… → Remove all…'),
          ('My edits disappeared after opening a file', 'It was opened as a different workbook: switch back in the selector'),
          ('Saving was stopped', 'Read the reason; usually no empty row is left. Use New version in template format'),
          ('Save back is greyed out', 'Open the file with Open workbook… (not by dragging) in Chrome or Edge'),
          ('The import created a second publisher', 'The unique name didn’t match. Generate again with the exact name'),
          ('Windows refuses to run import.ps1', 'Unblock-File .\\import.ps1, or -ExecutionPolicy Bypass for that session'),
          ('Can I share the HTML file?', 'Yes: it holds only the empty template. Nothing is uploaded')],
      0.54, 1.65, 12.2, (4.6, 7.6), size=13, h=4.9)
notes(s, 'The user guide’s troubleshooting section has the full list.')

# ---------------------------------------------------------------- close
s = new(13, keep=(13, 18, 19, 16, 17))
put(s.placeholders[13], 'Quality and resources')
put(s.placeholders[18], 'How it is tested')
put(s.placeholders[16], ['241 automated browser tests run on every change.', 'Generated solutions are unpacked and re-packed with Microsoft’s own packaging tool.',
                         'The solution XML was compared element by element with solutions exported by Dataverse.',
                         'import.ps1 is parsed with PowerShell 7.4 and run against a stand-in for pac.'], size=15)
put(s.placeholders[19], 'Where to find more')
put(s.placeholders[17], ['User guide with marked screenshots (Markdown and Word)', 'A narrated 2-minute walkthrough video',
                         'The short overview deck', 'Import_Test_Model.xlsx: three tables to try the whole flow',
                         'All in the SharePointPCF repository, under data-model-explorer'], size=15)
bullets(s.placeholders[16]); bullets(s.placeholders[17])
notes(s, 'Everything referred to in this deck is in the repository under data-model-explorer: the HTML file, docs/User_Guide.md and .docx, docs/Walkthrough.mp4, docs/Data_Model_Explorer_Overview.pptx, and samples/Import_Test_Model.xlsx.')

s = new(2, keep=(14,))
ph = s.placeholders[14]; ph.left, ph.top, ph.width, ph.height = Inches(0.6), Inches(2.1), Inches(12.1), Inches(1.2)
put(ph, 'Getting started')
text(s, 0.6, 3.6, 12.1, 0.5, ['Four steps from workbook to Dataverse'], size=20, color=MUTED, align=PP_ALIGN.CENTER, name='Subtitle')
for i, (n, b) in enumerate([('1', 'Open the Data Model Explorer HTML file in Chrome or Edge'), ('2', 'Drop the data-model workbook on it, or a metadata export'),
                            ('3', 'Work through Checks until there are no errors'), ('4', 'Build the solution and run import.ps1 against development')]):
    x = 0.6 + i * 2.97
    card(s, x, 4.35, 2.8, 1.45, name='Step')
    text(s, x + 0.25, 4.45, 0.6, 0.45, [n], size=24, color=PINK, bold=True)
    text(s, x + 0.25, 4.9, 2.35, 0.8, [b], size=14, color=DARK)
notes(s, 'Start with the import test model to see the whole flow end to end before using a real model.')

prs.save(OUT); print('saved', OUT, len(prs.slides), 'slides')
